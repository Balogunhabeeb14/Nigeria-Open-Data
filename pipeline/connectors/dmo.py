"""Connector for the Debt Management Office (DMO): what each state owes.

The DMO publishes, every quarter, "States and FCT Domestic Debt Stock as at
<date>": one document with a line per state and the Federal Capital Territory
in naira. The documents are listed at

    https://www.dmo.gov.ng/debt-profile/sub-national-debts

and each is downloaded from its page address plus '/file'.

This connector reads that list, downloads every quarter's document and turns
them into one long table (state × quarter). There is no API, so two things are
done to stay dependable:

* Files are parsed by content, not layout. Whatever the format (PDF, Excel or
  a web page), each line is searched for a state's name followed by an
  amount. A document that does not yield nearly all 37 lines is rejected and
  reported instead of being loaded half-read.
* Past quarters are kept in the landing folder and not downloaded again; only
  the newest few are re-fetched, because the DMO marks recent figures
  provisional and may revise them.
"""
from __future__ import annotations

import io
import logging
import re
import time
from datetime import datetime
from urllib.parse import urljoin

import pandas as pd

from pipeline.connectors.base import BaseConnector
from pipeline.geo import STATE_SPELLINGS, STATES

logger = logging.getLogger(__name__)

INDEX_URL = "https://www.dmo.gov.ng/debt-profile/sub-national-debts"
_PAGE_SIZE = 100
_MAX_PAGES = 10
_REFRESH_NEWEST = 3          # newest documents are always downloaded again
_MIN_STATES = 34             # of 37; fewer means the document was not read properly

_TITLE = re.compile(r"domestic\s+debt.*?as\s+at\s+(?P<date>[A-Za-z]+\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4})", re.IGNORECASE)
_SLUG = re.compile(r"domestic-debt", re.IGNORECASE)
_AMOUNT = re.compile(r"(?<![\w.])\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![\w])|(?<![\w.,])\d{7,}(?:\.\d+)?(?![\w,])")
_QUARTER_ENDS = {3: 1, 6: 2, 9: 3, 12: 4}


class DMOParseError(RuntimeError):
    """A document did not contain a readable state-by-state table."""


# ── Reading the list of documents ─────────────────────────────────────────────

def parse_report_date(text: str) -> datetime | None:
    """'...as at March 31st, 2026' → datetime(2026, 3, 31). None if no date."""
    match = _TITLE.search(text or "")
    if not match:
        return None
    cleaned = re.sub(r"(\d)(st|nd|rd|th)", r"\1", match.group("date"), flags=re.IGNORECASE).replace(",", " ")
    for fmt in ("%B %d %Y", "%b %d %Y"):
        try:
            return datetime.strptime(" ".join(cleaned.split()), fmt)
        except ValueError:
            continue
    return None


def quarter_of(date: datetime) -> str:
    """datetime(2026, 3, 31) → '2026Q1'."""
    return f"{date.year}Q{(date.month - 1) // 3 + 1}"


def parse_index(html: str, base_url: str = INDEX_URL) -> list[dict]:
    """Find the domestic-debt documents on one page of the DMO's list.

    Returns [{title, url, file_url, date, period}], newest first. Documents
    about external debt, and anything without a quarter-end date, are left out.
    """
    from bs4 import BeautifulSoup

    found: dict[str, dict] = {}
    for link in BeautifulSoup(html, "html.parser").find_all("a", href=True):
        title = " ".join(link.get_text(" ", strip=True).split())
        if "external" in title.lower():
            continue
        date = parse_report_date(title)
        if not date or date.month not in _QUARTER_ENDS:
            continue
        url = urljoin(base_url, link["href"]).split("?")[0].rstrip("/")
        if url.endswith("/file"):
            url = url[: -len("/file")]
        period = quarter_of(date)
        # One document per quarter; keep the first (newest listing) seen.
        found.setdefault(period, {
            "title": title, "url": url, "file_url": url + "/file",
            "date": date.date().isoformat(), "period": period,
        })
    return sorted(found.values(), key=lambda d: d["period"], reverse=True)


# ── Reading one document ──────────────────────────────────────────────────────

def document_text(content: bytes) -> str:
    """Plain text of a PDF, Excel workbook or web page, one row per line."""
    if content[:5] == b"%PDF-":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if content[:2] == b"PK":
        from openpyxl import load_workbook
        lines = []
        for sheet in load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = []
                for cell in row:
                    if cell is None:
                        continue
                    cells.append(f"{cell:,.2f}" if isinstance(cell, (int, float)) and abs(cell) >= 1e6 else str(cell))
                if cells:
                    lines.append("  ".join(cells))
        return "\n".join(lines)
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content.decode("utf-8", errors="replace"), "html.parser")
    rows = [" ".join(tr.get_text(" ", strip=True).split()) for tr in soup.find_all("tr")]
    return "\n".join(rows) if rows else soup.get_text("\n")


def parse_state_amounts(text: str) -> tuple[dict[str, float], float | None]:
    """Read {state code: amount} and the stated total from a document's text.

    A line counts when it names a state (in any of its usual spellings) and
    carries an amount after the name; the last amount on the line is taken,
    which is the debt stock in the DMO's single-column layout.
    """
    amounts: dict[str, float] = {}
    total: float | None = None
    for raw in text.splitlines():
        line = " ".join(raw.split())
        if not line:
            continue
        numbers = _AMOUNT.findall(line)
        if not numbers:
            continue
        value = float(numbers[-1].replace(",", ""))
        if re.match(r"^(grand\s+)?total\b", line, re.IGNORECASE):
            total = value
            continue
        upper = line.upper()
        for spelling, code in STATE_SPELLINGS:
            match = re.search(rf"(?<![A-Z]){re.escape(spelling.upper())}(?![A-Z])", upper)
            if match and code not in amounts:
                # The amount must come after the name, not be a serial number before it.
                if _AMOUNT.search(line[match.end():]):
                    amounts[code] = value
                break
    return amounts, total


# ── Connector ─────────────────────────────────────────────────────────────────

class DMOConnector(BaseConnector):
    """State-by-state domestic debt, every published quarter."""

    source_id = "DMO"

    def discover(self) -> list[dict]:
        """Every quarterly domestic-debt document the DMO lists, newest first."""
        documents: dict[str, dict] = {}
        for page in range(_MAX_PAGES):
            url = f"{INDEX_URL}?limit={_PAGE_SIZE}&limitstart={page * _PAGE_SIZE}"
            fresh = [d for d in parse_index(self._http_get(url).text) if d["period"] not in documents]
            if not fresh:
                break
            documents.update({d["period"]: d for d in fresh})
            time.sleep(float(self.options.get("request_delay_seconds") or 1))
        ordered = sorted(documents.values(), key=lambda d: d["period"], reverse=True)
        logger.info("DMO: %d quarterly domestic-debt documents listed", len(ordered))
        return ordered

    def read_document(self, document: dict, refresh: bool = False) -> tuple[dict[str, float], float | None]:
        """Download (or reuse) one document and return its state amounts and total."""
        path = self.landing_path(f"states-domestic-debt-{document['period']}.bin")
        if refresh or not path.exists() or path.stat().st_size == 0:
            self._download_file(document["file_url"], path)
            time.sleep(float(self.options.get("request_delay_seconds") or 1))
        amounts, total = parse_state_amounts(document_text(path.read_bytes()))
        if len(amounts) < _MIN_STATES:
            path.unlink(missing_ok=True)       # do not keep a file we could not read
            raise DMOParseError(
                f"{document['title']}: only {len(amounts)} of {len(STATES)} states could be read"
            )
        return amounts, total

    def fetch(self, dataset_config: dict, since: str | None = None) -> pd.DataFrame:
        """One row per state and quarter (kind 'states'), or per quarter (kind 'total').

        Columns: REF_AREA, TIME_PERIOD, OBS_VALUE, SOURCE_URL, NOTE.
        A document that cannot be read is skipped with a warning; the others
        still load.
        """
        kind = dataset_config.get("kind", "states")
        rows: list[dict] = []
        unreadable: list[str] = []

        for position, document in enumerate(self.discover()):
            try:
                amounts, total = self.read_document(document, refresh=position < _REFRESH_NEWEST)
            except Exception as exc:
                unreadable.append(document["period"])
                logger.warning("DMO: skipping %s (%s)", document["period"], exc)
                continue

            computed = sum(amounts.values())
            note = None
            missing = sorted(STATES[c] for c in set(STATES) - set(amounts))
            if missing:
                note = f"Not found in the source document for this quarter: {', '.join(missing)}."
                logger.warning("DMO %s: %s", document["period"], note)
            elif total and abs(computed - total) > 0.005 * total:
                note = f"States read sum to {computed:,.2f}; the document's total line says {total:,.2f}."
                logger.warning("DMO %s: %s", document["period"], note)

            if kind == "total":
                rows.append({
                    "REF_AREA": "NG", "TIME_PERIOD": document["period"],
                    # The DMO's own total where it prints one and all states were read.
                    "OBS_VALUE": total if (total and len(amounts) == len(STATES) and not note) else computed,
                    "SOURCE_URL": document["url"], "NOTE": note,
                })
            else:
                rows.extend(
                    {"REF_AREA": code, "TIME_PERIOD": document["period"], "OBS_VALUE": value,
                     "SOURCE_URL": document["url"], "NOTE": None}
                    for code, value in sorted(amounts.items())
                )

        if unreadable:
            logger.warning("DMO: %d document(s) could not be read: %s", len(unreadable), ", ".join(unreadable))
        df = pd.DataFrame(rows, columns=["REF_AREA", "TIME_PERIOD", "OBS_VALUE", "SOURCE_URL", "NOTE"])
        return df.sort_values(["TIME_PERIOD", "REF_AREA"]).reset_index(drop=True) if not df.empty else df
